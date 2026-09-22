from __future__ import annotations

import io
import json
import os
import re
import uuid
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterable

import bleach
import markdown
import requests
from flask import current_app
from sqlalchemy import text
from werkzeug.utils import secure_filename

from .models import Attachment, Note, NoteRevision, Notebook, Tag, db


ALLOWED_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".gif", ".webp",
    ".pdf", ".docx", ".xlsx", ".xlsm", ".pptx",
    ".txt", ".md", ".csv", ".json", ".xml", ".log",
}

SKILLS = {
    "summarise": {
        "label": "Summarise",
        "mode": "answer",
        "prompt": "Summarise the note accurately and concisely. Preserve key facts, decisions and numbers.",
    },
    "rewrite": {
        "label": "Improve wording",
        "mode": "content",
        "prompt": "Rewrite the note for clarity and structure while preserving its meaning. Return only the revised Markdown note.",
    },
    "actions": {
        "label": "Extract actions",
        "mode": "answer",
        "prompt": "Extract decisions, actions, owners and dates from this note. Use concise Markdown checklists. Do not invent missing owners or dates.",
    },
    "tags": {
        "label": "Suggest tags",
        "mode": "tags",
        "prompt": "Suggest 3 to 8 short reusable tags for this note. Return only a comma-separated list of tags.",
    },
    "title": {
        "label": "Suggest title",
        "mode": "title",
        "prompt": "Suggest one concise descriptive title for this note. Return only the title, without quotes.",
    },
    "mermaid": {
        "label": "Create diagram",
        "mode": "answer",
        "prompt": "Create a useful Mermaid diagram based only on the note. Return a fenced mermaid code block plus one short explanation.",
    },
    "related": {
        "label": "Find related",
        "mode": "answer",
        "prompt": "Identify the most closely related notes in the supplied context and explain the relationship briefly. Refer to note titles exactly.",
    },
}


def ensure_runtime_dirs() -> None:
    Path(current_app.instance_path).mkdir(parents=True, exist_ok=True)
    Path(current_app.config["UPLOAD_FOLDER"]).mkdir(parents=True, exist_ok=True)


def migrate_database() -> None:
    """Add v2 fields non-destructively to the old v1 SQLite database."""
    db.create_all()
    engine = db.engine
    with engine.begin() as conn:
        columns = {row[1] for row in conn.execute(text("PRAGMA table_info(note)"))}
        additions = {
            "notebook_id": "INTEGER",
            "created_at": "DATETIME",
            "updated_at": "DATETIME",
            "pinned": "BOOLEAN NOT NULL DEFAULT 0",
            "favourite": "BOOLEAN NOT NULL DEFAULT 0",
            "archived": "BOOLEAN NOT NULL DEFAULT 0",
            "source_url": "VARCHAR(1000)",
            "deleted_at": "DATETIME",
            "metadata_json": "TEXT NOT NULL DEFAULT '{}'",
        }
        for name, ddl in additions.items():
            if name not in columns:
                conn.execute(text(f"ALTER TABLE note ADD COLUMN {name} {ddl}"))
        conn.execute(text("UPDATE note SET created_at = COALESCE(created_at, CURRENT_TIMESTAMP)"))
        conn.execute(text("UPDATE note SET updated_at = COALESCE(updated_at, created_at, CURRENT_TIMESTAMP)"))

        try:
            conn.execute(text("CREATE VIRTUAL TABLE IF NOT EXISTS note_fts USING fts5(title, content, content='note', content_rowid='id')"))
            conn.execute(text("""
                CREATE TRIGGER IF NOT EXISTS note_ai AFTER INSERT ON note BEGIN
                    INSERT INTO note_fts(rowid, title, content) VALUES (new.id, new.title, new.content);
                END
            """))
            conn.execute(text("""
                CREATE TRIGGER IF NOT EXISTS note_ad AFTER DELETE ON note BEGIN
                    INSERT INTO note_fts(note_fts, rowid, title, content) VALUES ('delete', old.id, old.title, old.content);
                END
            """))
            conn.execute(text("""
                CREATE TRIGGER IF NOT EXISTS note_au AFTER UPDATE OF title, content ON note BEGIN
                    INSERT INTO note_fts(note_fts, rowid, title, content) VALUES ('delete', old.id, old.title, old.content);
                    INSERT INTO note_fts(rowid, title, content) VALUES (new.id, new.title, new.content);
                END
            """))
            conn.execute(text("INSERT INTO note_fts(note_fts) VALUES ('rebuild')"))
            current_app.config["FTS5_AVAILABLE"] = True
        except Exception:
            current_app.config["FTS5_AVAILABLE"] = False

    inbox = Notebook.query.filter_by(name="Inbox").first()
    if inbox is None:
        inbox = Notebook(name="Inbox")
        db.session.add(inbox)
        db.session.commit()
    Note.query.filter(Note.notebook_id.is_(None)).update({Note.notebook_id: inbox.id}, synchronize_session=False)
    db.session.commit()


def note_to_dict(note: Note, include_content: bool = True) -> dict:
    payload = {
        "id": note.id,
        "title": note.title,
        "notebook_id": note.notebook_id,
        "notebook": note.notebook.name if note.notebook else None,
        "tags": [tag.name for tag in sorted(note.tags, key=lambda t: t.name.lower())],
        "pinned": bool(note.pinned),
        "favourite": bool(note.favourite),
        "archived": bool(note.archived),
        "deleted": note.deleted_at is not None,
        "source_url": note.source_url,
        "created_at": note.created_at.isoformat() if note.created_at else None,
        "updated_at": note.updated_at.isoformat() if note.updated_at else None,
        "legacy_image": note.image,
        "attachments": [
            {
                "id": a.id,
                "name": a.original_name,
                "mime_type": a.mime_type,
                "size_bytes": a.size_bytes,
                "has_text": bool(a.extracted_text.strip()),
            }
            for a in sorted(note.attachments, key=lambda x: x.created_at)
        ],
    }
    if include_content:
        payload["content"] = note.content or ""
    else:
        payload["preview"] = re.sub(r"\s+", " ", note.content or "").strip()[:180]
    return payload


def get_or_create_tags(names: Iterable[str]) -> list[Tag]:
    clean = []
    seen = set()
    for raw in names:
        name = re.sub(r"\s+", " ", str(raw or "")).strip()[:80]
        key = name.lower()
        if name and key not in seen:
            clean.append(name)
            seen.add(key)
    tags: list[Tag] = []
    for name in clean:
        tag = Tag.query.filter(db.func.lower(Tag.name) == name.lower()).first()
        if not tag:
            tag = Tag(name=name)
            db.session.add(tag)
        tags.append(tag)
    return tags


def snapshot_note(note: Note, force: bool = False) -> None:
    latest = NoteRevision.query.filter_by(note_id=note.id).order_by(NoteRevision.created_at.desc()).first()
    if not force and latest and datetime.utcnow() - latest.created_at < timedelta(minutes=5):
        return
    db.session.add(NoteRevision(
        note_id=note.id,
        title=note.title,
        content=note.content or "",
        notebook_id=note.notebook_id,
        tags_csv=", ".join(tag.name for tag in note.tags),
    ))


def search_notes(query: str = "", notebook_id: int | None = None, view: str = "all", limit: int = 100) -> list[Note]:
    q = (query or "").strip()
    base = Note.query
    if view == "trash":
        base = base.filter(Note.deleted_at.isnot(None))
    else:
        base = base.filter(Note.deleted_at.is_(None))
        if view == "archive":
            base = base.filter(Note.archived.is_(True))
        else:
            base = base.filter(Note.archived.is_(False))
        if view == "favourites":
            base = base.filter(Note.favourite.is_(True))
        if view == "pinned":
            base = base.filter(Note.pinned.is_(True))
    if notebook_id:
        base = base.filter(Note.notebook_id == notebook_id)

    if not q:
        return base.order_by(Note.pinned.desc(), Note.updated_at.desc()).limit(limit).all()

    ids: list[int] = []
    if current_app.config.get("FTS5_AVAILABLE"):
        tokens = [t for t in re.findall(r"[\w-]+", q) if t]
        if tokens:
            match = " AND ".join(f'"{t.replace(chr(34), "")}"*' for t in tokens)
            try:
                rows = db.session.execute(
                    text("SELECT rowid FROM note_fts WHERE note_fts MATCH :match ORDER BY rank LIMIT :limit"),
                    {"match": match, "limit": limit},
                ).all()
                ids = [int(row[0]) for row in rows]
            except Exception:
                ids = []
    if ids:
        order = {note_id: idx for idx, note_id in enumerate(ids)}
        notes = base.filter(Note.id.in_(ids)).all()
        notes.sort(key=lambda n: order.get(n.id, 999999))
        return notes[:limit]

    like = f"%{q}%"
    notes = base.filter(db.or_(Note.title.ilike(like), Note.content.ilike(like))).order_by(Note.updated_at.desc()).limit(limit).all()
    if len(notes) < limit:
        attachment_note_ids = [
            r[0]
            for r in db.session.query(Attachment.note_id).filter(Attachment.extracted_text.ilike(like)).limit(limit).all()
        ]
        extras = base.filter(Note.id.in_(attachment_note_ids)).order_by(Note.updated_at.desc()).all() if attachment_note_ids else []
        present = {n.id for n in notes}
        notes.extend(n for n in extras if n.id not in present)
    return notes[:limit]


def render_markdown(source: str) -> str:
    html = markdown.markdown(
        source or "",
        extensions=["fenced_code", "tables", "sane_lists", "nl2br"],
    )
    allowed_tags = set(bleach.sanitizer.ALLOWED_TAGS).union({
        "p", "pre", "code", "h1", "h2", "h3", "h4", "h5", "h6", "table", "thead", "tbody", "tr", "th", "td",
        "hr", "br", "blockquote", "ul", "ol", "li", "input", "img", "span", "div",
    })
    attrs = {
        "a": ["href", "title", "target", "rel"],
        "code": ["class"],
        "pre": ["class"],
        "input": ["type", "checked", "disabled"],
        "img": ["src", "alt", "title"],
        "div": ["class"],
        "span": ["class"],
    }
    return bleach.clean(
        html,
        tags=allowed_tags,
        attributes=attrs,
        protocols={"http", "https", "mailto", "data"},
        strip=True,
    )


def _extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            import fitz
            with fitz.open(path) as doc:
                return "\n\n".join(page.get_text("text") for page in doc)
        if ext == ".docx":
            from docx import Document
            doc = Document(path)
            chunks = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    chunks.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(chunks)
        if ext in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            chunks = []
            for ws in wb.worksheets:
                chunks.append(f"# {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    values = ["" if v is None else str(v) for v in row]
                    if any(values):
                        chunks.append(" | ".join(values))
            return "\n".join(chunks)
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(path)
            chunks = []
            for idx, slide in enumerate(prs.slides, start=1):
                chunks.append(f"# Slide {idx}")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        chunks.append(shape.text)
            return "\n".join(chunks)
        if ext in {".txt", ".md", ".csv", ".json", ".xml", ".log"}:
            return path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        return f"[Extraction failed: {exc}]"
    return ""


def save_attachment(file_storage, note: Note) -> Attachment:
    original = secure_filename(file_storage.filename or "attachment") or "attachment"
    ext = Path(original).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported attachment type: {ext or 'no extension'}")
    stored = f"{uuid.uuid4().hex}{ext}"
    target = Path(current_app.config["UPLOAD_FOLDER"]) / stored
    file_storage.save(target)
    attachment = Attachment(
        note_id=note.id,
        original_name=original,
        stored_name=stored,
        mime_type=file_storage.mimetype,
        size_bytes=target.stat().st_size,
        extracted_text=_extract_text(target)[:2_000_000],
    )
    db.session.add(attachment)
    db.session.commit()
    return attachment


def delete_attachment_file(attachment: Attachment) -> None:
    path = Path(current_app.config["UPLOAD_FOLDER"]) / attachment.stored_name
    if path.exists():
        path.unlink()


def settings_path() -> Path:
    return Path(current_app.instance_path) / "settings.json"


def load_ai_settings() -> dict:
    defaults = {
        "ollama_url": os.getenv("NOTES_OLLAMA_URL", "http://192.168.1.249:11434"),
        "model": os.getenv("NOTES_OLLAMA_MODEL", "qwen3:14b"),
        "embedding_model": os.getenv("NOTES_OLLAMA_EMBEDDING_MODEL", "nomic-embed-text"),
        "timeout_seconds": 120,
    }
    path = settings_path()
    if path.exists():
        try:
            saved = json.loads(path.read_text(encoding="utf-8"))
            defaults.update({k: saved[k] for k in defaults if k in saved})
        except Exception:
            pass
    return defaults


def save_ai_settings(settings: dict) -> dict:
    current = load_ai_settings()
    url = str(settings.get("ollama_url", current["ollama_url"])).strip().rstrip("/")
    if not url.startswith(("http://", "https://")):
        raise ValueError("Ollama URL must start with http:// or https://")
    current.update({
        "ollama_url": url,
        "model": str(settings.get("model", current["model"])).strip(),
        "embedding_model": str(settings.get("embedding_model", current["embedding_model"])).strip(),
        "timeout_seconds": max(5, min(600, int(settings.get("timeout_seconds", current["timeout_seconds"])))),
    })
    settings_path().write_text(json.dumps(current, indent=2), encoding="utf-8")
    return current


def ollama_models() -> list[str]:
    settings = load_ai_settings()
    response = requests.get(f"{settings['ollama_url']}/api/tags", timeout=10)
    response.raise_for_status()
    return [item.get("name") for item in response.json().get("models", []) if item.get("name")]


def ollama_generate(prompt: str, system: str = "") -> str:
    settings = load_ai_settings()
    payload = {
        "model": settings["model"],
        "prompt": prompt,
        "system": system,
        "stream": False,
        "options": {"temperature": 0.2},
    }
    response = requests.post(
        f"{settings['ollama_url']}/api/generate",
        json=payload,
        timeout=settings["timeout_seconds"],
    )
    response.raise_for_status()
    return str(response.json().get("response", "")).strip()


def context_for_ai(note: Note | None, question: str = "", scope: str = "current") -> str:
    chunks: list[str] = []
    if note:
        chunks.append(f"## Current note: {note.title}\n{(note.content or '')[:18000]}")
        for attachment in note.attachments:
            if attachment.extracted_text.strip():
                chunks.append(f"### Attachment: {attachment.original_name}\n{attachment.extracted_text[:5000]}")
    if scope == "all":
        matches = search_notes(question, limit=12) if question.strip() else search_notes(limit=12)
        for other in matches:
            if note and other.id == note.id:
                continue
            chunks.append(f"## Note: {other.title}\n{(other.content or '')[:4000]}")
    text_value = "\n\n".join(chunks)
    return text_value[:45000]


def run_skill(note: Note, skill_name: str, instruction: str = "") -> dict:
    skill = SKILLS.get(skill_name)
    if not skill:
        raise ValueError("Unknown skill")
    related_context = context_for_ai(
        note,
        instruction or note.title,
        "all" if skill_name == "related" else "current",
    )
    prompt = f"{skill['prompt']}\n\n{instruction.strip()}\n\nCONTEXT:\n{related_context}".strip()
    result = ollama_generate(
        prompt,
        system="You are the local Notes assistant. Use only the supplied notes and attachments. Never invent facts.",
    )
    mode = skill["mode"]
    if mode == "tags":
        tags = [re.sub(r"^[#\-*\s]+", "", t).strip() for t in re.split(r"[,\n]", result)]
        tags = [t for t in tags if t][:8]
        return {"kind": "tags", "skill": skill_name, "value": tags, "raw": result}
    if mode == "title":
        return {"kind": "title", "skill": skill_name, "value": result.strip().strip('"').splitlines()[0][:200]}
    if mode == "content":
        value = result
        fenced = re.match(r"^```(?:markdown)?\s*(.*?)\s*```$", result, flags=re.S | re.I)
        if fenced:
            value = fenced.group(1)
        return {"kind": "content", "skill": skill_name, "value": value}
    return {"kind": "answer", "skill": skill_name, "value": result}


def build_backup_zip() -> io.BytesIO:
    buffer = io.BytesIO()
    notes = Note.query.order_by(Note.id).all()
    payload = {
        "format": "notes-v2-backup",
        "created_at": datetime.utcnow().isoformat() + "Z",
        "notebooks": [{"id": n.id, "name": n.name} for n in Notebook.query.order_by(Notebook.id).all()],
        "notes": [note_to_dict(n, include_content=True) for n in notes],
        "revisions": [
            {
                "note_id": r.note_id,
                "title": r.title,
                "content": r.content,
                "notebook_id": r.notebook_id,
                "tags_csv": r.tags_csv,
                "created_at": r.created_at.isoformat(),
            }
            for r in NoteRevision.query.order_by(NoteRevision.id).all()
        ],
    }
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("notes.json", json.dumps(payload, indent=2, ensure_ascii=False))
        upload_root = Path(current_app.config["UPLOAD_FOLDER"])
        for attachment in Attachment.query.order_by(Attachment.id).all():
            path = upload_root / attachment.stored_name
            if path.exists():
                archive.write(path, f"attachments/{attachment.stored_name}")
    buffer.seek(0)
    return buffer
